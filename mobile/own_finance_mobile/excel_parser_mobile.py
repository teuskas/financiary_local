"""Excel parsing utilities shared by the mobile APK app."""

from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import load_workbook

MESI = [
    "GENNAIO",
    "FEBBRAIO",
    "MARZO",
    "APRILE",
    "MAGGIO",
    "GIUGNO",
    "LUGLIO",
    "AGOSTO",
    "SETTEMBRE",
    "OTTOBRE",
    "NOVEMBRE",
    "DICEMBRE",
]

TITOLI = {
    "investimenti": "INVESTIMENTI PER PIATTAFORMA",
    "guadagni": "GUADAGNI PER PIATTAFORMA",
    "inv_guad": "INV+GUAD PER PIATTAFORMA",
}


def _to_float(value: Any) -> float:
    if value is None:
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)

    text = str(value).strip()
    if not text:
        return 0.0

    normalized = text.replace(" ", "")
    if "," in normalized and "." in normalized:
        normalized = normalized.replace(".", "").replace(",", ".")
    else:
        normalized = normalized.replace(",", ".")

    try:
        return float(normalized)
    except ValueError:
        return 0.0


def _find_row_by_title(ws, title: str) -> int:
    for row_idx in range(1, ws.max_row + 1):
        cell = ws.cell(row=row_idx, column=1).value
        if str(cell).strip().upper() == title:
            return row_idx
    raise ValueError(f"Titolo non trovato: {title}")


def detect_current_year_sheet_from_bytes(xlsx_content: bytes) -> str:
    wb = load_workbook(BytesIO(xlsx_content), data_only=True, read_only=True)
    years = []
    for name in wb.sheetnames:
        text = str(name).strip()
        try:
            year = int(text)
        except ValueError:
            continue
        if 2000 <= year <= 2100:
            years.append(year)

    if not years:
        raise ValueError("Nessun foglio annuale trovato (es. 2026, 2027)")

    return str(max(years))


def extract_tables_from_bytes(xlsx_content: bytes, sheet_name: str) -> dict[str, list[dict[str, Any]]]:
    wb = load_workbook(BytesIO(xlsx_content), data_only=True, read_only=True)
    ws = wb[sheet_name]

    out: dict[str, list[dict[str, Any]]] = {}
    for key, title in TITOLI.items():
        out[key] = _extract_table(ws, title)
    return out


def _extract_table(ws, title: str) -> list[dict[str, Any]]:
    header_row = _find_row_by_title(ws, title)

    col_map: dict[str, int] = {}
    for col in range(2, min(ws.max_column, 20) + 1):
        val = ws.cell(row=header_row, column=col).value
        text = str(val).strip().upper()
        if text in MESI:
            col_map[text] = col
        elif text == "TOTALE":
            col_map["TOTALE"] = col

    rows: list[dict[str, Any]] = []
    for row in range(header_row + 1, ws.max_row + 1):
        label_raw = ws.cell(row=row, column=1).value
        label = str(label_raw).strip()

        if not label:
            continue

        if label.upper() in TITOLI.values() and label.upper() != title:
            break

        item: dict[str, Any] = {"Piattaforma": label}
        for mese, col in col_map.items():
            item[mese] = _to_float(ws.cell(row=row, column=col).value)
        rows.append(item)

        if label.upper() == "SOMMA":
            break

    return rows


def format_currency_it(value: float) -> str:
    base = f"{float(value):,.2f}"
    return base.replace(",", "#").replace(".", ",").replace("#", ".")


def table_to_text(table_rows: list[dict[str, Any]], max_columns: int = 6) -> str:
    if not table_rows:
        return "Nessun dato disponibile"

    all_cols = list(table_rows[0].keys())
    cols = all_cols[:max_columns]

    lines = [" | ".join(cols)]
    lines.append("-" * min(110, len(lines[0]) + 10))

    for row in table_rows:
        values = []
        for col in cols:
            val = row.get(col, "")
            if col != "Piattaforma":
                values.append(format_currency_it(val))
            else:
                values.append(str(val))
        lines.append(" | ".join(values))

    return "\n".join(lines)

