"""Smoke test del parser mobile senza accesso Dropbox."""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook

from own_finance_mobile.excel_parser_mobile import (
    detect_current_year_sheet_from_bytes,
    extract_tables_from_bytes,
)


def _make_sample_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "2026"

    # Header tabella investimenti
    ws.cell(row=1, column=1, value="INVESTIMENTI PER PIATTAFORMA")
    ws.cell(row=1, column=2, value="GENNAIO")
    ws.cell(row=1, column=3, value="FEBBRAIO")
    ws.cell(row=1, column=4, value="TOTALE")
    ws.cell(row=2, column=1, value="BONDORA")
    ws.cell(row=2, column=2, value=100)
    ws.cell(row=2, column=3, value=50)
    ws.cell(row=2, column=4, value=150)
    ws.cell(row=3, column=1, value="SOMMA")
    ws.cell(row=3, column=2, value=100)
    ws.cell(row=3, column=3, value=50)
    ws.cell(row=3, column=4, value=150)

    # Header tabella guadagni
    ws.cell(row=5, column=1, value="GUADAGNI PER PIATTAFORMA")
    ws.cell(row=5, column=2, value="GENNAIO")
    ws.cell(row=5, column=3, value="FEBBRAIO")
    ws.cell(row=5, column=4, value="TOTALE")
    ws.cell(row=6, column=1, value="BONDORA")
    ws.cell(row=6, column=2, value=10)
    ws.cell(row=6, column=3, value=5)
    ws.cell(row=6, column=4, value=15)
    ws.cell(row=7, column=1, value="SOMMA")
    ws.cell(row=7, column=2, value=10)
    ws.cell(row=7, column=3, value=5)
    ws.cell(row=7, column=4, value=15)

    # Header tabella inv+guad
    ws.cell(row=9, column=1, value="INV+GUAD PER PIATTAFORMA")
    ws.cell(row=9, column=2, value="GENNAIO")
    ws.cell(row=9, column=3, value="FEBBRAIO")
    ws.cell(row=9, column=4, value="TOTALE")
    ws.cell(row=10, column=1, value="BONDORA")
    ws.cell(row=10, column=2, value=110)
    ws.cell(row=10, column=3, value=55)
    ws.cell(row=10, column=4, value=165)
    ws.cell(row=11, column=1, value="SOMMA")
    ws.cell(row=11, column=2, value=110)
    ws.cell(row=11, column=3, value=55)
    ws.cell(row=11, column=4, value=165)

    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()


def main() -> None:
    content = _make_sample_xlsx()
    year = detect_current_year_sheet_from_bytes(content)
    assert year == "2026"

    tables = extract_tables_from_bytes(content, year)
    assert set(tables.keys()) == {"investimenti", "guadagni", "inv_guad"}
    assert tables["guadagni"][0]["GENNAIO"] == 10.0
    assert tables["inv_guad"][0]["TOTALE"] == 165.0

    print("[OK] smoke test parser mobile")


if __name__ == "__main__":
    main()

